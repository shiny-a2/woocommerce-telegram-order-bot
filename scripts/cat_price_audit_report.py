"""cat_price_audit_report.py — گزارشِ اکسلِ کاملِ وضعیتِ قیمتِ کت + فیلیپ‌پلین.

فرقش با اکسلِ روزانهٔ سینک: آن فقط «چه چیزی عوض شد» را می‌گوید، این «وضعیتِ همهٔ ۶۷۲ محصول» را
می‌گوید — از جمله آن‌هایی که عمداً دست نخوردند و علتش. برای وقتی که مالک می‌خواهد خودش همه‌چیز
را یک‌جا چک کند.

فقط خواندنی: هیچ چیزی روی سایت نوشته نمی‌شود.

    python scripts/cat_price_audit_report.py [--out data/cat-price-audit-<ts>.xlsx]
"""
from __future__ import annotations

import argparse
import asyncio
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import woo  # noqa: E402
import wt_pricesync as ps  # noqa: E402

MONEY = "#,##0"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
OK_FILL = PatternFill("solid", fgColor="E2EFDA")      # سبزِ کم‌رنگ — عوض شد
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")    # زردِ کم‌رنگ — نیاز به نگاهِ آدم
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")    # خاکستری — کاری لازم نیست

# وضعیت‌ها به فارسی، به‌ترتیبِ اهمیت برای مالک
S_AMBIGUOUS = "مبهم — خانواده چندقیمتی، دست نخورد"
S_NOPRICE = "بدونِ قیمت در کانال — دست نخورد"
S_EXACT = "رفرنسِ دقیق — قیمت از کانال"
S_FAMILY = "هم‌خانواده — قیمت از خطِ خانواده"
S_NOREF = "بدونِ رفرنس"


def _sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 30
    return ws


def _money(ws, row, cols):
    for c in cols:
        ws.cell(row=row, column=c).number_format = MONEY


async def collect():
    channel, meta = ps.load_channel_file()
    products = await ps.fetch_brand_products(woo)

    fam: dict[str, set] = {}
    for ref, price in channel.items():
        fk = ps.family_key(ref)
        if fk:
            fam.setdefault(fk, set()).add(price)

    rows = []
    matched = set()
    for p in products:
        ref = p["ref"]
        fk = ps.family_key(ref) if ref else None
        chan_price = None
        status = S_NOREF
        note = ""
        if not ref:
            pass
        elif ref in channel:
            matched.add(ref)
            chan_price = channel[ref]
            status = S_EXACT
        elif fk in fam and len(fam[fk]) == 1:
            chan_price = next(iter(fam[fk]))
            status = S_FAMILY
        elif fk in fam:
            status = S_AMBIGUOUS
            note = " / ".join(f"{x:,}" for x in sorted(fam[fk]))
        else:
            status = S_NOPRICE
            note = f"خانوادهٔ «{fk}» در کانال نیست" if fk else ""

        diff = None
        if chan_price is not None:
            diff = p["regular_price"] - chan_price
        rows.append({
            "id": p["id"], "name": p["name"], "ref": ref, "brand": p["brand"],
            "family": fk or "—", "site_price": p["regular_price"], "chan_price": chan_price,
            "diff": diff, "status": status, "note": note,
            "stock": "موجود" if p["stock_status"] == "instock" else "ناموجود",
            "qty": p["stock_quantity"] if p["manage_stock"] else None,
            "managed": "عددی" if p["manage_stock"] else "شرکتی",
        })

    orphan = sorted(r for r in channel if r not in matched)
    return rows, channel, fam, orphan, meta


def build(rows, channel, fam, orphan, meta, out_path):
    wb = Workbook()

    # ---------------- خلاصه ----------------
    ws = _sheet(wb, "خلاصه", ["مورد", "تعداد", "توضیح"], [46, 12, 72], first=True)
    by = {}
    for r in rows:
        by[r["status"]] = by.get(r["status"], 0) + 1
    mismatch = [r for r in rows if r["diff"] not in (None, 0)]
    summary = [
        ("کلِ محصولاتِ کت + فیلیپ‌پلین", len(rows), "برندهای مجاز روی سایت"),
        ("رفرنسِ منتشرشده در کانال", len(channel), "کانالِ تأمین‌کننده"),
        ("", "", ""),
        (S_EXACT, by.get(S_EXACT, 0), "قیمت مستقیم از کانال گرفته می‌شود"),
        (S_FAMILY, by.get(S_FAMILY, 0), "خطِ خانواده تک‌قیمت بود"),
        (S_AMBIGUOUS, by.get(S_AMBIGUOUS, 0), "قاعدهٔ خودت: خانوادهٔ چندقیمتی دست نمی‌خورد"),
        (S_NOPRICE, by.get(S_NOPRICE, 0), "تأمین‌کننده برای این‌ها قیمتی نداده"),
        ("", "", ""),
        ("قیمتِ سایت ≠ قیمتِ کانال (الان)", len(mismatch),
         "اگر بعد از سینک صفر نیست، یعنی جای بررسی دارد"),
        ("رفرنسِ کانال بدونِ محصول روی سایت", len(orphan), "کالایی که نداریم"),
    ]
    for i, (a, b, c) in enumerate(summary, start=2):
        ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b)
        ws.cell(row=i, column=3, value=c)
        if a in (S_AMBIGUOUS, S_NOPRICE):
            for col in (1, 2, 3):
                ws.cell(row=i, column=col).fill = WARN_FILL
    ws.cell(row=len(summary) + 3, column=1,
            value=f"ساخته‌شده: {datetime.now():%Y-%m-%d %H:%M}  ·  اسنپ‌شاتِ کانال: {meta.get('count')} رفرنس")

    # ---------------- همهٔ محصولات ----------------
    heads = ["شناسه", "نام محصول", "رفرنس", "برند", "خانواده", "قیمتِ سایت", "قیمتِ کانال",
             "اختلاف", "وضعیت", "توضیح", "موجودی", "نوع", "تعداد"]
    ws = _sheet(wb, "همهٔ محصولات", heads, [10, 52, 20, 16, 10, 16, 16, 14, 34, 30, 11, 10, 9])
    order = {S_AMBIGUOUS: 0, S_NOPRICE: 1, S_FAMILY: 2, S_EXACT: 3, S_NOREF: 4}
    for i, r in enumerate(sorted(rows, key=lambda x: (order.get(x["status"], 9),
                                                      x["family"], x["ref"])), start=2):
        vals = [r["id"], r["name"], r["ref"], r["brand"], r["family"], r["site_price"],
                r["chan_price"], r["diff"], r["status"], r["note"], r["stock"],
                r["managed"], r["qty"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
        _money(ws, i, (6, 7, 8))
        if r["status"] == S_AMBIGUOUS:
            fill = WARN_FILL
        elif r["status"] == S_NOPRICE:
            fill = GREY_FILL
        elif r["diff"] == 0:
            fill = OK_FILL
        else:
            fill = WARN_FILL
        for j in range(1, len(heads) + 1):
            ws.cell(row=i, column=j).fill = fill

    # ---------------- قیمت هنوز نمی‌خوانَد ----------------
    ws = _sheet(wb, "قیمت هنوز نمی‌خواند", heads, [10, 52, 20, 16, 10, 16, 16, 14, 34, 30, 11, 10, 9])
    for i, r in enumerate(sorted(mismatch, key=lambda x: -abs(x["diff"] or 0)), start=2):
        vals = [r["id"], r["name"], r["ref"], r["brand"], r["family"], r["site_price"],
                r["chan_price"], r["diff"], r["status"], r["note"], r["stock"],
                r["managed"], r["qty"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=i, column=j, value=v)
        _money(ws, i, (6, 7, 8))
    if not mismatch:
        ws.cell(row=2, column=1, value="✅ هیچ اختلافی نیست — قیمتِ همهٔ محصولاتِ دارای قیمتِ کانال دقیقاً برابر است.")

    # ---------------- مبهم ----------------
    ws = _sheet(wb, "مبهم (تصمیمِ تو)",
                ["خانواده", "شناسه", "نام محصول", "رفرنس", "قیمتِ فعلیِ سایت",
                 "قیمت‌های کانال در این خانواده", "موجودی"],
                [10, 10, 52, 20, 18, 42, 11])
    amb = [r for r in rows if r["status"] == S_AMBIGUOUS]
    for i, r in enumerate(sorted(amb, key=lambda x: (x["family"], x["ref"])), start=2):
        for j, v in enumerate([r["family"], r["id"], r["name"], r["ref"],
                               r["site_price"], r["note"], r["stock"]], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.fill = WARN_FILL
        _money(ws, i, (5,))

    # ---------------- بدونِ قیمت در کانال ----------------
    ws = _sheet(wb, "بدونِ قیمت در کانال",
                ["خانواده", "تعدادِ محصول", "نمونه رفرنس‌ها"], [12, 14, 90])
    nop = [r for r in rows if r["status"] == S_NOPRICE]
    groups: dict[str, list] = {}
    for r in nop:
        groups.setdefault(r["family"], []).append(r["ref"])
    for i, (fk, refs) in enumerate(sorted(groups.items(), key=lambda x: -len(x[1])), start=2):
        ws.cell(row=i, column=1, value=fk)
        ws.cell(row=i, column=2, value=len(refs))
        ws.cell(row=i, column=3, value="، ".join(sorted(refs)[:12]) + ("، …" if len(refs) > 12 else ""))

    # ---------------- رفرنسِ کانال بدونِ محصول ----------------
    ws = _sheet(wb, "رفرنسِ کانال بی‌محصول",
                ["رفرنسِ کانال", "قیمتِ کانال", "خانواده", "نکته"], [24, 18, 12, 46])
    site_refs = {r["ref"] for r in rows if r["ref"]}
    for i, ref in enumerate(orphan, start=2):
        fk = ps.family_key(ref)
        # رفرنس‌های عادی چهاربخشی‌اند؛ بخشِ اضافه معمولاً یعنی قیمت/سایز چسبیده به رفرنس
        odd = ref.count(".") >= 4
        ws.cell(row=i, column=1, value=ref)
        ws.cell(row=i, column=2, value=channel[ref]).number_format = MONEY
        ws.cell(row=i, column=3, value=fk)
        ws.cell(row=i, column=4,
                value="⚠️ رفرنسِ نامتعارف (بخشِ اضافه) — احتمالاً بد خوانده شده و خانواده را مبهم می‌کند"
                if odd else ("این کالا روی سایت نیست" if fk not in {r['family'] for r in rows}
                             else "روی سایت محصولِ این رفرنس نیست"))
        if odd:
            for j in range(1, 5):
                ws.cell(row=i, column=j).fill = WARN_FILL

    wb.save(out_path)
    return out_path, len(rows), len(mismatch), len(amb), len(nop), len(orphan)


async def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="")
    args = ap.parse_args()
    out = args.out or os.path.join("data", f"cat-price-audit-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    rows, channel, fam, orphan, meta = await collect()
    path, n, mism, amb, nop, orp = build(rows, channel, fam, orphan, meta, out)
    print(f"محصولات: {n} · اختلافِ قیمت: {mism} · مبهم: {amb} · بی‌قیمت: {nop} · رفرنسِ بی‌محصول: {orp}")
    print(f"اکسل: {path}")


if __name__ == "__main__":
    asyncio.run(main())
