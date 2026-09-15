"""irantimer_brand_catalog.py — کاتالوگِ کاملِ برندهای irantimer + تطبیق با برندهای سایتِ خودمان.

چرا لازم شد: جدولِ `BRANDS` در irantimer_extract_job.py دستی بود و فقط ۲۳ ورودی داشت، پس
`/brand سیکو` جواب نمی‌داد در حالی که سیکو روی irantimer هست (Brand=43، ۴۰۱۵ محصول).
کاتالوگِ قدیمیِ data/irantimer_brands.json هم نام‌هایش خراب بود («لیست محصولات»).

خروجی:
  data/irantimer_brands_v2.json  — [{id, slug, name, total}] برای همهٔ برندها
  data/irantimer_brand_match.xlsx — تطبیقِ برندهای irantimer با ترم‌های برندِ سایت

فقط خواندنی؛ هیچ چیزی روی سایت نوشته نمی‌شود.

    python scripts/irantimer_brand_catalog.py [--limit N] [--no-excel]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import irantimer as it  # noqa: E402
import woo  # noqa: E402

BRAND_ATTR_ID = 103          # pa_نام-برند
OUT_JSON = os.path.join("data", "irantimer_brands_v2.json")
OUT_XLSX = os.path.join("data", "irantimer_brand_match.xlsx")

_ZWNJ = "‌"


def norm_fa(s: str) -> str:
    """نرمال‌سازیِ نامِ فارسی برای تطبیق: فاصله/نیم‌فاصله حذف، ی/ک عربی → فارسی."""
    s = (s or "").strip()
    s = s.replace("ي", "ی").replace("ك", "ک").replace("ۀ", "ه").replace("أ", "ا").replace("إ", "ا")
    s = s.replace(_ZWNJ, "").replace("‏", "").replace("‎", "")
    s = re.sub(r"[\s\-_.]+", "", s)
    return s


def site_brand_terms() -> list[dict]:
    """همهٔ ترم‌های اتریبیوتِ «نام برند» روی سایت → [{id, name, count}]."""
    out, page = [], 1
    while page <= 30:
        rows = woo._get_sync(f"products/attributes/{BRAND_ATTR_ID}/terms",
                             {"per_page": 100, "page": page, "_fields": "id,name,count"})
        if not rows:
            break
        out.extend(rows)
        if len(rows) < 100:
            break
        page += 1
    return out


def build_catalog(limit: int = 0, pause: float = 0.25) -> list[dict]:
    brands = it.list_brands()
    print(f"لینکِ برند در صفحهٔ رسمی: {len(brands)}")
    rows = []
    for i, b in enumerate(brands, start=1):
        if limit and i > limit:
            break
        if b["id"] == 0:                      # «ساعت مچی» = خودِ گروه، برند نیست
            continue
        try:
            info = it.brand_info(b["id"])
        except Exception as e:  # noqa: BLE001
            print(f"  [{i}/{len(brands)}] {b['slug']}: خطا {type(e).__name__}")
            continue
        # نامِ h1 معتبرتر از alt است؛ alt فقط برای برندهای اسلایدر درست بود
        name = info["name"] or b["name"]
        rows.append({"id": b["id"], "slug": b["slug"], "name": name,
                     "title": info["title"], "total": info["total"]})
        if i % 20 == 0:
            print(f"  … {i}/{len(brands)}")
        time.sleep(pause)
    rows.sort(key=lambda r: -r["total"])
    return rows


def match(rows: list[dict], terms: list[dict]) -> tuple[list, list, list]:
    by_norm = {}
    for t in terms:
        by_norm.setdefault(norm_fa(t["name"]), t)
    matched, it_only = [], []
    used = set()
    for r in rows:
        t = by_norm.get(norm_fa(r["name"]))
        if t:
            used.add(t["id"])
            matched.append({**r, "term_id": t["id"], "term_name": t["name"],
                            "site_count": t.get("count", 0)})
        else:
            it_only.append(r)
    site_only = [t for t in terms if t["id"] not in used and (t.get("count") or 0) > 0]
    site_only.sort(key=lambda t: -(t.get("count") or 0))
    return matched, it_only, site_only


def to_excel(matched, it_only, site_only, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    HDR_FILL = PatternFill("solid", fgColor="1F3864")
    HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
    wb = Workbook()

    def sheet(title, headers, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        ws.row_dimensions[1].height = 28
        return ws

    ws = sheet("روی هر دو (قابلِ استخراج)",
               ["برند", "Brand ID", "محصول در کاتالوگِ منبع", "ترمِ سایت", "محصول روی سایتِ ما", "اسلاگ"],
               [24, 11, 20, 12, 20, 28], first=True)
    for i, r in enumerate(sorted(matched, key=lambda x: -x["total"]), start=2):
        for j, v in enumerate([r["name"], r["id"], r["total"], r["term_id"],
                               r["site_count"], r["slug"]], start=1):
            ws.cell(row=i, column=j, value=v)

    ws = sheet("فقط کاتالوگِ منبع (برندِ تازه)",
               ["برند", "Brand ID", "محصول در کاتالوگِ منبع", "اسلاگ"], [24, 11, 20, 30])
    for i, r in enumerate(sorted(it_only, key=lambda x: -x["total"]), start=2):
        for j, v in enumerate([r["name"], r["id"], r["total"], r["slug"]], start=1):
            ws.cell(row=i, column=j, value=v)

    ws = sheet("فقط سایتِ ما", ["برند", "ترمِ سایت", "محصول روی سایتِ ما"], [26, 12, 20])
    for i, t in enumerate(site_only, start=2):
        for j, v in enumerate([t["name"], t["id"], t.get("count", 0)], start=1):
            ws.cell(row=i, column=j, value=v)

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--no-excel", action="store_true")
    args = ap.parse_args()

    rows = build_catalog(args.limit)
    json.dump(rows, open(OUT_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"\nکاتالوگ: {len(rows)} برند → {OUT_JSON}")

    terms = site_brand_terms()
    print(f"ترمِ برند روی سایتِ ما: {len(terms)}")
    matched, it_only, site_only = match(rows, terms)
    print(f"روی هر دو: {len(matched)} · فقط کاتالوگِ منبع: {len(it_only)} · فقط سایتِ ما: {len(site_only)}")
    print(f"جمعِ محصولاتِ قابلِ استخراج (برندهای مشترک): {sum(r['total'] for r in matched):,}")

    print("\nبزرگ‌ترین برندهای مشترک:")
    for r in sorted(matched, key=lambda x: -x["total"])[:15]:
        print(f"  {r['name']:<18} کاتالوگِ منبع {r['total']:>6,}   سایتِ ما {r['site_count']:>6,}")

    if not args.no_excel:
        to_excel(matched, it_only, site_only, OUT_XLSX)
        print(f"\nاکسل: {OUT_XLSX}")


if __name__ == "__main__":
    main()
