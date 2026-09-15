"""irantimer_learn.py — یادگیری از اصلاحاتِ اپراتور روی اکسلِ استخراج.

جریانِ کار: ما اکسل می‌فرستیم → اپراتور تطبیق می‌دهد و سلول‌ها را اصلاح می‌کند → فایل را برمی‌گرداند
→ این ماژول «اکسلِ اصلاح‌شده» را با «اکسلی که خودمان ساختیم» مقایسه می‌کند و از تفاوت‌ها قاعده می‌سازد،
تا دفعهٔ بعد همان اشتباه تکرار نشود.

اصلِ کار: نگاشت روی **مقدارِ خامِ منبع** یاد گرفته می‌شود، نه روی خروجیِ غلطِ خودمان. اگر فقط
«خروجیِ غلط → مقدارِ درست» را ذخیره کنیم، دو مقدارِ خامِ متفاوت که تصادفاً یک خروجیِ غلط داده‌اند
هر دو با یک قاعده عوض می‌شوند و این غلط است. مقدارِ خام از صفحهٔ محصولِ کاتالوگِ منبع بازیابی می‌شود
(شناسهٔ محصول از ستونِ «لینک» همان اکسل درمی‌آید).

سه نوع چیز یاد گرفته می‌شود:
  • value_maps  — «مقدارِ خامِ کاتالوگِ منبع» → «ترمِ سایت» برای ویژگی‌هایی که منبعِ مستقیم دارند (SRC).
                  برندمستقل است و روی همهٔ برندهای بعدی اثر می‌گذارد.
  • brand_const — ثابت‌های هر برند (گارانتی، گارانتی‌کننده، نگین) که مَپر عمداً خالی می‌گذارد تا
                  اپراتور پر کند. اگر در یک بسته اکثر سطرها یک مقدار گرفتند، ثابتِ آن برند می‌شود.
  • per_ref     — اصلاحِ تک‌محصولی، برای ویژگی‌های قاعده‌محور (طرح صفحه/استایل/…) که نمی‌شود از یک
                  نمونه قاعدهٔ کلی ساخت. اینجا حدس نمی‌زنیم؛ فقط همان محصول را درست می‌کنیم و
                  موردش را برای بازبینیِ قاعده علامت می‌زنیم.

تضادها (یک مقدارِ خام که دو بار دو جورِ مختلف اصلاح شده) هرگز بی‌صدا حل نمی‌شوند: آخرین نظرِ
اپراتور اعمال می‌شود ولی مورد در `conflicts` ثبت و در گزارش برجسته می‌شود.

هیچ AI ای در این ماژول نیست و هیچ چیزی روی سایت نوشته نمی‌شود.
"""
from __future__ import annotations

import datetime as _dt
import json
import os
import re

import openpyxl

import irantimer as it
import irantimer_map as im

_HERE = os.path.dirname(os.path.abspath(__file__))
_DATA = os.path.join(_HERE, "data")
OVERRIDES = os.path.join(_DATA, "atefeh_overrides.json")
_DETAILS = os.path.join(_DATA, "irantimer_details")

REF_COL = "رفرانس"
LINK_COL = "لینک"
BRAND_COL = "نام برند"

# ویژگی‌هایی که مَپر عمداً خالی می‌گذارد تا اپراتور پر کند → ثابتِ برند
BRAND_CONST_ATTRS = {"گارانتی", "گارانتی کننده در ایران", "نگین", "موارد گارانتی"}
# نسبتِ لازم برای اینکه یک مقدار «ثابتِ برند» شناخته شود
BRAND_CONST_RATIO = 0.6
_PID_RE = re.compile(r"/Fa-Product-(\d+)/")


def _now() -> str:
    return _dt.datetime.now().isoformat(timespec="seconds")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "").strip())


# ---------------------------------------------------------------- overrides ----
def load_overrides() -> dict:
    try:
        with open(OVERRIDES, encoding="utf-8") as f:
            d = json.load(f)
    except Exception:  # noqa: BLE001 — نبودنِ فایل حالتِ عادیِ اولین اجراست
        d = {}
    d.setdefault("version", 1)
    for k in ("value_maps", "brand_const", "per_ref"):
        d.setdefault(k, {})
    for k in ("conflicts", "log"):
        d.setdefault(k, [])
    return d


def save_overrides(d: dict) -> None:
    d["updated"] = _now()
    os.makedirs(_DATA, exist_ok=True)
    tmp = OVERRIDES + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=1)
    os.replace(tmp, OVERRIDES)


# ------------------------------------------------------------------- excel ----
def read_sheet(path: str) -> dict[str, dict]:
    """اکسلِ استخراج → {رفرانس: {ستون: مقدار}}. ستون‌ها با نامِ سرصفحه خوانده می‌شوند،
    پس جابه‌جا شدنِ ستون‌ها یا حذفِ ستونِ عکس توسطِ اپراتور مشکلی درست نمی‌کند."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [_norm(h) for h in next(rows)]
    except StopIteration:
        return {}
    idx = {h: i for i, h in enumerate(header) if h}
    if REF_COL not in idx:
        raise ValueError(f"ستونِ «{REF_COL}» در {os.path.basename(path)} نیست.")
    out = {}
    for r in rows:
        ref = _norm(r[idx[REF_COL]] if idx[REF_COL] < len(r) else "")
        if not ref:
            continue
        out[ref.upper()] = {h: _norm(r[i]) if i < len(r) else "" for h, i in idx.items()}
    return out


def diff(original: dict, corrected: dict, attrs: list[str] | None = None) -> list[dict]:
    """تفاوت‌های سلولی. فقط ستون‌های خروجیِ سایت مقایسه می‌شوند (ستون‌های مرجع/عکس نه)."""
    attrs = attrs or im.OUT_ATTRS
    changes = []
    for ref, corr in corrected.items():
        orig = original.get(ref)
        if orig is None:
            continue                      # سطری که در اکسلِ ما نبوده — بازبینیِ دستی، یادگیری نه
        for a in attrs:
            if a in (REF_COL, BRAND_COL):
                continue
            before, after = _norm(orig.get(a, "")), _norm(corr.get(a, ""))
            if before == after:
                continue
            changes.append({"ref": ref, "attr": a, "before": before, "after": after,
                            "link": corr.get(LINK_COL) or orig.get(LINK_COL) or "",
                            "brand": corr.get(BRAND_COL) or orig.get(BRAND_COL) or ""})
    return changes


# ------------------------------------------------------------ source values ----
def _detail_cached(pid: str) -> dict | None:
    p = os.path.join(_DETAILS, f"{pid}.json")
    if os.path.exists(p):
        try:
            with open(p, encoding="utf-8") as f:
                return json.load(f)
        except Exception:  # noqa: BLE001
            pass
    try:
        d = it.parse_detail(str(pid))
    except Exception:  # noqa: BLE001 — صفحه ممکن است حذف شده باشد؛ بدونِ مقدارِ خام فقط per_ref یاد می‌گیریم
        return None
    os.makedirs(_DETAILS, exist_ok=True)
    try:
        with open(p, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False)
    except Exception:  # noqa: BLE001
        pass
    return d


def source_value(change: dict, detail: dict | None) -> str:
    """مقدارِ خامِ کاتالوگِ منبع که این ویژگی از آن ساخته شده (اگر منبعِ مستقیم داشته باشد)."""
    if not detail:
        return ""
    key = im.SRC.get(change["attr"])
    if not key:
        return ""
    return _norm((detail.get("specs") or {}).get(key, ""))


# ---------------------------------------------------------------- learning ----
def learn(changes: list[dict], batch_name: str = "", apply: bool = True) -> dict:
    """از تغییرات قاعده می‌سازد. apply=False یعنی فقط گزارش (چیزی ذخیره نمی‌شود)."""
    ov = load_overrides()
    stats = {"changes": len(changes), "value_map": 0, "per_ref": 0,
             "brand_const": 0, "conflicts": 0, "no_source": 0}
    details: dict[str, dict | None] = {}
    learned: list[dict] = []
    # نامزدهای ثابتِ برند: {(brand, attr): {value: count}}
    const_votes: dict[tuple, dict] = {}

    for ch in changes:
        pid = ""
        m = _PID_RE.search(ch.get("link") or "")
        if m:
            pid = m.group(1)
        if pid and pid not in details:
            details[pid] = _detail_cached(pid)
        det = details.get(pid)
        raw = source_value(ch, det)
        brand = _norm(ch.get("brand"))
        attr = ch["attr"]

        if attr in BRAND_CONST_ATTRS and brand:
            const_votes.setdefault((brand, attr), {})
            const_votes[(brand, attr)][ch["after"]] = const_votes[(brand, attr)].get(ch["after"], 0) + 1
            continue

        if raw:
            slot = ov["value_maps"].setdefault(attr, {})
            prev = slot.get(raw)
            if prev and _norm(prev.get("value")) != _norm(ch["after"]):
                ov["conflicts"].append({
                    "when": _now(), "attr": attr, "source": raw, "ref": ch["ref"],
                    "was": prev.get("value"), "now": ch["after"], "batch": batch_name,
                })
                stats["conflicts"] += 1
            if prev and _norm(prev.get("value")) == _norm(ch["after"]):
                prev["support"] = int(prev.get("support", 1)) + 1
                prev["updated"] = _now()
            else:
                slot[raw] = {"value": ch["after"], "support": 1, "updated": _now(),
                             "example_ref": ch["ref"]}
                stats["value_map"] += 1
                learned.append({"kind": "value_map", "attr": attr, "source": raw,
                                "value": ch["after"]})
        else:
            # ویژگیِ قاعده‌محور یا منبعِ ناموجود → فقط همین محصول، بدونِ تعمیم
            ov["per_ref"].setdefault(ch["ref"], {})[attr] = {
                "value": ch["after"], "updated": _now(), "was": ch["before"]}
            stats["per_ref"] += 1
            if not det:
                stats["no_source"] += 1
            learned.append({"kind": "per_ref", "attr": attr, "ref": ch["ref"],
                            "value": ch["after"]})

    # ثابت‌های برند: فقط وقتی اکثریتِ روشنی وجود دارد
    for (brand, attr), votes in const_votes.items():
        total = sum(votes.values())
        value, n = max(votes.items(), key=lambda kv: kv[1])
        if total >= 3 and n / total >= BRAND_CONST_RATIO:
            ov["brand_const"].setdefault(brand, {})[attr] = {
                "value": value, "support": n, "of": total, "updated": _now()}
            stats["brand_const"] += 1
            learned.append({"kind": "brand_const", "brand": brand, "attr": attr, "value": value})
        else:
            # اکثریت نبود → تک‌تک به per_ref می‌روند تا چیزی از دست نرود
            for ch in changes:
                if ch["attr"] == attr and _norm(ch.get("brand")) == brand:
                    ov["per_ref"].setdefault(ch["ref"], {})[attr] = {
                        "value": ch["after"], "updated": _now(), "was": ch["before"]}
                    stats["per_ref"] += 1

    ov["log"].append({"when": _now(), "batch": batch_name, **stats})
    if apply:
        save_overrides(ov)
    return {"stats": stats, "learned": learned,
            "conflicts": ov["conflicts"][-stats["conflicts"]:] if stats["conflicts"] else []}


def learn_from_files(original_path: str, corrected_path: str, apply: bool = True) -> dict:
    orig = read_sheet(original_path)
    corr = read_sheet(corrected_path)
    ch = diff(orig, corr)
    res = learn(ch, batch_name=os.path.basename(corrected_path), apply=apply)
    res["rows_original"] = len(orig)
    res["rows_corrected"] = len(corr)
    res["changes"] = ch
    return res


def find_original(corrected_path: str) -> str | None:
    """اکسلِ اصلی‌ای که این فایل اصلاحِ آن است را پیدا می‌کند.

    اول با نامِ فایل تلاش می‌شود (اپراتور معمولاً همان نام را نگه می‌دارد)، وگرنه آن فایلِ
    irantimer-*.xlsx که بیشترین رفرنسِ مشترک را با فایلِ اصلاح‌شده دارد.
    """
    base = os.path.basename(corrected_path)
    cand = os.path.join(_DATA, base)
    if os.path.exists(cand) and os.path.abspath(cand) != os.path.abspath(corrected_path):
        return cand
    try:
        corr_refs = set(read_sheet(corrected_path))
    except Exception:  # noqa: BLE001
        return None
    if not corr_refs:
        return None
    best, best_n = None, 0
    for f in os.listdir(_DATA):
        if not (f.startswith("irantimer-") and f.endswith(".xlsx")):
            continue
        p = os.path.join(_DATA, f)
        if os.path.abspath(p) == os.path.abspath(corrected_path):
            continue
        try:
            n = len(corr_refs & set(read_sheet(p)))
        except Exception:  # noqa: BLE001
            continue
        if n > best_n:
            best, best_n = p, n
    return best if best_n >= max(3, len(corr_refs) // 4) else None
